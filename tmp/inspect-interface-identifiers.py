from __future__ import annotations

from pathlib import Path

from docx import Document
from pypdf import PdfReader


DOCX_FILES = [
    Path(r"docs\参考资料\甘肃文件\附件3：甘肃省检查检验结果互认平台对接技术方案V1.0.docx"),
    Path(r"docs\参考资料\甘肃文件\附件2：甘肃省检查检验结果互认平台数据交换接口规范V1.0.docx"),
    Path(r"docs\参考资料\附件4：检验检查结果互认技术方案 (1).docx"),
]

PDF_FILES = [
    Path(r"docs\参考资料\DB64+2122-2025+医疗机构检查检验结果互认规范.pdf"),
    Path(r"docs\参考资料\DB35_T+2153-2023医疗机构检查检验结果互认共享数据传输及应用要求.pdf"),
]

TERMS = (
    "流水",
    "业务标识",
    "查询标识",
    "就诊流水",
    "医嘱",
    "申请单",
    "申请序号",
    "候选",
    "互认结果",
    "反馈",
    "采纳",
    "引用",
    "existsReCureInfo",
    "queryId",
    "serial",
    "order",
    "apply",
    "visit",
)


def clean(value: str) -> str:
    return " ".join(value.replace("\x00", "").split())


for path in DOCX_FILES:
    print(f"\n=== DOCX {path} ===")
    document = Document(path)
    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = clean(paragraph.text)
        if text and any(term.lower() in text.lower() for term in TERMS):
            print(f"P{index}: {text}")
    for table_index, table in enumerate(document.tables, start=1):
        rows = [" | ".join(clean(cell.text) for cell in row.cells) for row in table.rows]
        matching = [index for index, row in enumerate(rows) if any(term.lower() in row.lower() for term in TERMS)]
        if matching:
            print(f"-- TABLE {table_index} rows={len(rows)} matches={matching} --")
            for row_index, row in enumerate(rows):
                print(f"T{table_index}R{row_index + 1}: {row}")


for path in PDF_FILES:
    print(f"\n=== PDF {path} ===")
    reader = PdfReader(path)
    for page_index, page in enumerate(reader.pages, start=1):
        text = clean(page.extract_text() or "")
        if any(term.lower() in text.lower() for term in TERMS):
            snippets = []
            lowered = text.lower()
            for term in TERMS:
                start = lowered.find(term.lower())
                if start >= 0:
                    snippets.append(text[max(0, start - 180): start + 420])
            print(f"PAGE {page_index}: {' || '.join(dict.fromkeys(snippets))}")
